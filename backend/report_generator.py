"""
report_generator.py
────────────────────
Extracted from the old streamlit_app.py.
Generates professional PDF (ReportLab) and DOCX (python-docx) reports
from the report_agent text output + optional chart image.
"""

import io
import re
from datetime import datetime

from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer,
    HRFlowable, Image as RLImage,
)
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_JUSTIFY

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ── Brand colours ─────────────────────────────────────────────────────────────
BRAND_BLUE     = colors.HexColor("#1E88E5")
BRAND_CYAN     = colors.HexColor("#00D4FF")
DARK_TEXT       = colors.HexColor("#2c3e50")
MID_GRAY        = colors.HexColor("#7f8c8d")

LINE_SPACING    = 18
SPACE_BEFORE    = 6
SPACE_AFTER     = 6

SECTION_KEYWORDS = [
    "executive", "summary", "finding", "analysis", "conclusion",
    "performance", "trend", "market", "user", "subscription",
    "visual", "overview", "detail", "indicator",
]


# ── Helpers ───────────────────────────────────────────────────────────────────

def _normalise_report_text(raw):
    """Accept str / dict / list and return a clean string."""
    if isinstance(raw, dict):
        raw = raw.get("text", "")
    elif isinstance(raw, list):
        parts = []
        for item in raw:
            parts.append(item.get("text", str(item)) if isinstance(item, dict) else str(item))
        raw = "\n".join(parts)
    elif not isinstance(raw, str):
        raw = str(raw) if raw else ""

    raw = re.sub(r"\n+", "\n", raw)
    raw = re.sub(r"\\n", "\n", raw)
    raw = re.sub(r" +", " ", raw)
    return raw


def _clean_line(text: str) -> str:
    return text.replace("#", "").replace("**", "").replace("*", "").strip()


def _get_cleaned_lines(report_text: str) -> list[str]:
    return [
        _clean_line(line)
        for line in report_text.split("\n")
        if _clean_line(line)
        and _clean_line(line) not in ("---", "--", "-")
        and not _clean_line(line).lower().startswith(("date:", "generated:"))
    ]


def _is_heading(line: str) -> bool:
    is_numbered = len(line) > 2 and line[0].isdigit() and line[1] in ".)"
    is_keyword  = any(kw in line.lower() for kw in SECTION_KEYWORDS) and len(line) < 80
    return is_numbered or is_keyword


# ══════════════════════════════════════════════════════════════════════════════
# PDF Generation
# ══════════════════════════════════════════════════════════════════════════════

def generate_pdf(report_text: str, chart_bytes: io.BytesIO | None = None) -> io.BytesIO:
    """Return a BytesIO containing the finished PDF."""
    report_text = _normalise_report_text(report_text)
    cleaned = _get_cleaned_lines(report_text)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=letter,
        leftMargin=0.9 * inch, rightMargin=0.9 * inch,
        topMargin=1.0 * inch, bottomMargin=0.9 * inch,
    )

    style_title = ParagraphStyle(
        "MainTitle", fontSize=20, fontName="Helvetica-Bold",
        textColor=DARK_TEXT, alignment=TA_CENTER,
        spaceBefore=SPACE_BEFORE, spaceAfter=SPACE_AFTER, leading=28,
    )
    style_heading = ParagraphStyle(
        "SubHeading", fontSize=12, fontName="Helvetica-Bold",
        textColor=DARK_TEXT, alignment=TA_LEFT,
        spaceBefore=SPACE_BEFORE, spaceAfter=SPACE_AFTER, leading=int(12 * 1.5),
    )
    style_body = ParagraphStyle(
        "Body", fontSize=10, fontName="Helvetica",
        textColor=DARK_TEXT, alignment=TA_JUSTIFY,
        spaceBefore=SPACE_BEFORE, spaceAfter=SPACE_AFTER, leading=LINE_SPACING,
    )
    style_bullet = ParagraphStyle(
        "Bullet", fontSize=10, fontName="Helvetica",
        textColor=DARK_TEXT, leftIndent=18,
        spaceBefore=SPACE_BEFORE, spaceAfter=SPACE_AFTER, leading=LINE_SPACING,
    )
    style_footer = ParagraphStyle(
        "Footer", fontSize=8, fontName="Helvetica",
        textColor=MID_GRAY, alignment=TA_CENTER,
    )
    style_chart_heading = ParagraphStyle(
        "ChartHeading", fontSize=12, fontName="Helvetica-Bold",
        textColor=DARK_TEXT, spaceBefore=SPACE_BEFORE, spaceAfter=SPACE_AFTER,
    )

    story = []

    # Title
    story.append(Spacer(1, 0.2 * inch))
    title = cleaned[0] if cleaned else "Commodity Report"
    story.append(Paragraph(title, style_title))
    story.append(HRFlowable(width="100%", thickness=1, color=BRAND_CYAN, spaceAfter=14))
    story.append(Spacer(1, 0.1 * inch))

    # Body
    for line in cleaned[1:]:
        if not line:
            story.append(Spacer(1, 6))
        elif _is_heading(line):
            story.append(Paragraph(line, style_heading))
        elif line.startswith(("•", "-")):
            story.append(Paragraph(f"• {line.lstrip('•- ').strip()}", style_bullet))
        elif "|" in line and "---" not in line:
            continue
        else:
            story.append(Paragraph(line, style_body))

    # Chart
    if chart_bytes:
        story.append(Spacer(1, 0.2 * inch))
        story.append(HRFlowable(width="100%", thickness=1, color=BRAND_CYAN, spaceAfter=8))
        story.append(Paragraph("Visual Analysis", style_chart_heading))
        chart_bytes.seek(0)
        story.append(RLImage(chart_bytes, width=6.0 * inch, height=3.4 * inch))

    # Footer
    story.append(Spacer(1, 0.3 * inch))
    story.append(HRFlowable(width="100%", thickness=1, color=BRAND_BLUE, spaceAfter=6))
    story.append(Paragraph(
        "Confidential  •  Transgraph Commodity Risk Management  •  AI-Generated Report",
        style_footer,
    ))

    doc.build(story)
    buf.seek(0)
    return buf


# ══════════════════════════════════════════════════════════════════════════════
# DOCX Generation
# ══════════════════════════════════════════════════════════════════════════════

def _set_color(run, hex_color: str):
    run.font.color.rgb = RGBColor(
        int(hex_color[0:2], 16),
        int(hex_color[2:4], 16),
        int(hex_color[4:6], 16),
    )


def _set_line_spacing(para, space_before=6, space_after=6, line_spacing=1.5):
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.space_after  = Pt(space_after)
    para.paragraph_format.line_spacing = line_spacing


def _add_hr(doc, color_hex="00D4FF", thickness=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(4)
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(thickness))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color_hex)
    pBdr.append(bottom)
    pPr.append(pBdr)


def generate_docx(report_text: str, chart_bytes: io.BytesIO | None = None) -> io.BytesIO:
    """Return a BytesIO containing the finished DOCX."""
    report_text = _normalise_report_text(report_text)
    cleaned = _get_cleaned_lines(report_text)

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)

    # Title
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_line_spacing(title_para, 6, 6)
    tr = title_para.add_run(cleaned[0] if cleaned else "Report")
    tr.bold = True
    tr.font.size = Pt(20)
    _set_color(tr, "2c3e50")

    _add_hr(doc, "00D4FF", 6)
    doc.add_paragraph()

    # Body
    for line in cleaned[1:]:
        if not line:
            p = doc.add_paragraph()
            _set_line_spacing(p)
            continue

        if _is_heading(line):
            h = doc.add_paragraph()
            _set_line_spacing(h, 6, 6)
            run = h.add_run(line)
            run.bold = True
            run.font.size = Pt(11)
            _set_color(run, "2c3e50")

        elif line.startswith(("•", "-")):
            bp = doc.add_paragraph(style="List Bullet")
            _set_line_spacing(bp, 6, 6)
            bp.paragraph_format.left_indent = Inches(0.25)
            r = bp.add_run(line.lstrip("•- ").strip())
            r.font.size = Pt(10)
            _set_color(r, "2c3e50")

        elif "|" in line and "---" not in line:
            continue

        else:
            bp = doc.add_paragraph()
            _set_line_spacing(bp, 6, 6)
            r = bp.add_run(line)
            r.font.size = Pt(10)
            _set_color(r, "2c3e50")

    # Chart
    if chart_bytes:
        _add_hr(doc, "00D4FF", 6)
        ch = doc.add_paragraph()
        _set_line_spacing(ch, 6, 6)
        cr = ch.add_run("Visual Analysis")
        cr.bold = True
        cr.font.size = Pt(12)
        _set_color(cr, "2c3e50")
        chart_bytes.seek(0)
        doc.add_picture(chart_bytes, width=Inches(5.8))

    # Footer
    _add_hr(doc, "1E88E5", 10)
    fp = doc.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_line_spacing(fp, 4, 4)
    fr = fp.add_run(
        "Confidential  •  Transgraph Commodity Risk Management  •  AI-Generated Report"
    )
    fr.font.size = Pt(8)
    _set_color(fr, "7f8c8d")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
